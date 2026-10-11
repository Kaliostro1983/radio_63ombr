"""Донесення за результатами радіорозвідки (Форма 1.2.17).

Двоколонкова таблиця «Висновки | Радіоперехоплення», згрупована за напрямком
(один — Лиманський) та бригадно-полковими підрозділами (пули).

Ліва колонка — аналітичний висновок (частина НАД роздільником «🦁 63 ОМБр 🦁»,
у БД це `analytical_conclusions.conclusion_text`). Права — картка перехоплення
(привʼязане повідомлення: Дата/Частота/Локація/Підрозділ/Час/Хто/Кому/Текст).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


@dataclass
class RerItem:
    conclusion: str = ""   # ліва колонка (над роздільником)
    points: list = field(default_factory=list)  # точки (MGRS) події з висновку
    date: str = ""         # Дата
    time: str = ""         # Час
    freq: str = ""         # Частота
    unit: str = ""         # Підрозділ
    location: str = ""     # Локація (район)
    who: str = ""          # Хто (caller)
    whom: str = ""         # Кому (callee)
    text: str = ""         # Текст (тіло перехоплення)


_MGRS_RE = re.compile(r"^(\d{1,2}[A-Za-z])\s*([A-Za-z]{2})\s*(\d+)$")


def format_mgrs(raw: str | None) -> str:
    """'37UDQ2841529225' → '37U DQ 28415 29225' (парне число цифр → схід/північ)."""
    s = (raw or "").replace(" ", "").upper()
    m = _MGRS_RE.match(s)
    if not m:
        return (raw or "").strip()
    gz, sq, digits = m.group(1), m.group(2), m.group(3)
    if len(digits) % 2 == 0 and digits:
        h = len(digits) // 2
        return f"{gz} {sq} {digits[:h]} {digits[h:]}"
    return f"{gz} {sq} {digits}"


# ── Витяг бригадно-полкового підрозділу з опису радіомережі ─────────────────
# Порядок важливий: довші/специфічні токени раніше; дивізії (мсд/тд) НЕ беремо.
_REG_RE = re.compile(
    r"(\d+)\s*(омсбр|опбпс|обпс|мсбр|тбр|адбр|сабр|бр|осп|мсп|сп|тп)",
    re.IGNORECASE,
)
_DIV_RE = re.compile(r"(\d+)\s*(мсд|тд)", re.IGNORECASE)


def subunit_group(unit: str | None) -> str:
    """Ключ пулу: ярлик бригадно-полкового підрозділу, напр. '36 мсп', '164 омсбр'.

    Якщо полк/бригаду не знайдено — намагається дивізію, інакше 'Інше'.
    Використовується ЛИШЕ для групування/сортування (не для заголовка).
    """
    u = (unit or "").strip()
    if not u:
        return "Інше"
    m = _REG_RE.search(u)
    if m:
        return f"{m.group(1)} {m.group(2).lower()}"
    m = _DIV_RE.search(u)
    if m:
        return f"{m.group(1)} {m.group(2).lower()}"
    return "Інше"


def subunit_full(unit: str | None) -> str:
    """Повна назва підрозділу для заголовка пулу — від полку/бригади до кінця
    (відкидає початковий батальйонний/ротний/штабний префікс).

    'зрбатр 36 мсп 67 мсд 25 А' → '36 мсп 67 мсд 25 А';
    '3 мсб 31 мсп 67 мсд 25 ЗА' → '31 мсп 67 мсд 25 ЗА';
    '164 омсбр 25 ЗА' → '164 омсбр 25 ЗА'.
    """
    u = (unit or "").strip()
    if not u:
        return "Інше"
    m = _REG_RE.search(u)
    if m:
        return u[m.start():].strip()
    m = _DIV_RE.search(u)
    if m:
        return u[m.start():].strip()
    return u


def _group_sort_key(label: str) -> tuple:
    """Сортування пулів: спершу за номером, потім за типом."""
    m = re.match(r"(\d+)\s*(.*)", label)
    if m:
        return (0, int(m.group(1)), m.group(2))
    return (1, 0, label)


# ── Верстка docx ────────────────────────────────────────────────────────────
def _set_base_font(doc: Document) -> None:
    doc.styles["Normal"].font.size = Pt(11)


def _center(cell) -> None:
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _bold_cell(cell, text: str, *, size: int = 11, center: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)


def _intercept_cell_text(it: RerItem):
    """Рядки правої колонки (мітка, значення)."""
    return [
        ("Дата: ", it.date),
        ("Частота: ", it.freq),
        ("Локація: ", it.location),
        ("Підрозділ: ", it.unit),
        ("Час: ", it.time),
        ("Хто: ", it.who),
        ("Кому: ", it.whom),
        ("Текст:", ""),
    ]


def _fill_conclusion_cell(cell, it: RerItem) -> None:
    """Ліва колонка: текст висновку + нумерований перелік точок (MGRS) події."""
    cell.text = ""
    conc = (it.conclusion or "").strip()
    pts = list(it.points or [])
    if pts:
        # текст завершується двокрапкою «…в р-ні точки:» — далі перелік точок.
        if conc and not conc.endswith(":"):
            conc = conc.rstrip(".") + ":"
        cell.paragraphs[0].add_run(conc).font.size = Pt(11)
        for i, p in enumerate(pts, 1):
            para = cell.add_paragraph()
            para.add_run(f"{i}. {p}").font.size = Pt(11)
    else:
        # без точок — прибрати «висячу» двокрапку, завершити крапкою.
        if conc.endswith(":"):
            conc = conc[:-1].rstrip() + "."
        elif conc and conc[-1] not in ".!?":
            conc = conc + "."
        cell.paragraphs[0].add_run(conc).font.size = Pt(11)


def _fill_intercept_cell(cell, it: RerItem) -> None:
    cell.text = ""
    first = True
    for label, value in _intercept_cell_text(it):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(11)
        if value:
            rv = p.add_run(value)
            rv.font.size = Pt(11)
    # Текст перехоплення — окремими рядками під міткою «Текст:».
    body = (it.text or "").strip()
    if body:
        for line in body.splitlines():
            line = line.strip()
            if not line:
                continue
            p = cell.add_paragraph()
            p.add_run(line).font.size = Pt(11)


def build_rer_donesennya_docx_bytes(
    *,
    direction: str,
    pools: list[tuple[str, list[RerItem]]],
    report_dt: datetime,
    period_label: str,
    filename_range: str,
) -> tuple[bytes, str]:
    """Зібрати docx. `pools` — [(повна назва підрозділу, [RerItem…])…] у порядку.

    Повертає (bytes, filename). Ім'я: «Форма 1.2.17 DD.MM.YYYY (HH:MM-HH:MM).docx».
    """
    date_str = report_dt.strftime("%d.%m.%Y")
    dt_str = report_dt.strftime("%d.%m.%Y %H:%M")

    doc = Document()
    _set_base_font(doc)

    # Шапка
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0.add_run("Форма №1.2.17/УВС").font.size = Pt(11)

    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run("Донесення за результатами ведення радіоелектронної розвідки")
    r1.bold = True
    r1.font.size = Pt(14)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run(f"(станом на {dt_str}; період: {period_label})").font.size = Pt(12)

    doc.add_paragraph("")

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"

    # Заголовок таблиці
    hdr = table.add_row().cells
    _bold_cell(hdr[0], "Висновки", center=True)
    _bold_cell(hdr[1], "Радіоперехоплення", center=True)

    # Напрямок (обʼєднаний рядок)
    dir_row = table.add_row().cells
    dir_cell = dir_row[0].merge(dir_row[1])
    _bold_cell(dir_cell, direction, size=12, center=True)

    total = 0
    for label, items in pools:
        if not items:
            continue
        su_row = table.add_row().cells
        su_cell = su_row[0].merge(su_row[1])
        _bold_cell(su_cell, label)
        for it in items:
            total += 1
            row = table.add_row().cells
            _fill_conclusion_cell(row[0], it)
            _fill_intercept_cell(row[1], it)

    if total == 0:
        empty_row = table.add_row().cells
        ec = empty_row[0].merge(empty_row[1])
        ec.text = "За вказаний період аналітичних висновків не зафіксовано."

    filename = f"Форма 1.2.17 {date_str} ({filename_range}).docx"
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue(), filename


__all__ = [
    "RerItem",
    "subunit_group",
    "subunit_full",
    "format_mgrs",
    "build_rer_donesennya_docx_bytes",
]
